from docx import Document
from zhiyu.parser.document_parser import DocumentParser

def test_table_order(tmp_path):
    path = tmp_path / "a.docx"
    source = Document()
    source.add_paragraph("before")
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "单元格"
    table.cell(0, 1).text = "table text"
    source.add_paragraph("after")
    source.save(path)
    text = DocumentParser().parse(path).text
    assert text.index("before") < text.index("单元格") < text.index("table text") < text.index("after")

